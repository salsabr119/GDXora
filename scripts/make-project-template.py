# Builds public/templates/project-import.xlsx — the Excel template read by the
# project importer (src/screens/projectImport.js). Keys in the hidden column E
# of «المشروع» and the column order of «جدول الكميات» are the contract.
#   python3 scripts/make-project-template.py public/templates/project-import.xlsx
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

out = sys.argv[1]
GOLD, BLACK, REQ, GREY = "FFCE08", "111111", "FDECEC", "F6F6F6"
thin = Side(style="thin", color="D9D9D9"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
hfont = Font(bold=True, color="FFFFFF", name="Arial", size=11); hfill = PatternFill("solid", fgColor=BLACK)
base = Font(name="Arial", size=11); bold = Font(name="Arial", size=11, bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)

wb = Workbook()

# ── Sheet 1: project header ─────────────────────────────────────────────
ws = wb.active; ws.title = "المشروع"; ws.sheet_view.rightToLeft = True
ws["A1"] = "بيانات المشروع · Project details"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
ws["A2"] = "الحقول المظللة بالأحمر إلزامية. اكتب القيم في العمود C فقط."; ws["A2"].font = Font(name="Arial", size=10, color="B00020")
for c, h in zip("ABCD", ["الحقل", "Field", "القيمة · Value", "ملاحظة · Note"]):
    ws[f"{c}4"] = h; ws[f"{c}4"].font = hfont; ws[f"{c}4"].fill = hfill; ws[f"{c}4"].alignment = center
fields = [  # key, ar, en, required, note, kind
    ("code", "كود المشروع", "Project code", False, "اختياري، فريد داخل الشركة. مثال: PRJ-001", "text"),
    ("name_ar", "اسم المشروع (عربي)", "Project name (AR)", True, "", "text"),
    ("name_en", "اسم المشروع (إنجليزي)", "Project name (EN)", False, "", "text"),
    ("customer", "العميل", "Customer", True, "اسم العميل كما هو في النظام، أو رقمه الضريبي. إذا العميل جديد عبّئ ورقة «العميل»", "text"),
    ("contract_no", "رقم العقد", "Contract no.", False, "", "text"),
    ("contract_value", "قيمة العقد قبل الضريبة (ر.س)", "Contract value excl. VAT", True, "يفترض أن تساوي إجمالي جدول الكميات", "money"),
    ("retention_pct", "نسبة محتجزات الضمان %", "Retention %", False, "من 0 إلى 100. مثال: 10", "pct"),
    ("advance_amount", "الدفعة المقدمة المستلمة (ر.س)", "Advance payment received", False, "", "money"),
    ("advance_recovery_pct", "نسبة استرداد الدفعة المقدمة من كل مستخلص %", "Advance recovery %", False, "من 0 إلى 100", "pct"),
    ("opening_advance_recovered", "الدفعة المقدمة المستردة سابقاً (ر.س)", "Advance already recovered", False, "للمشاريع القائمة فقط: المسترد في مستخلصات قبل النظام", "money"),
    ("opening_billing_count", "عدد المستخلصات الصادرة قبل النظام", "Billings issued before the system", False, "للمشاريع القائمة فقط. المستخلص القادم يأخذ الرقم التالي", "int"),
    ("start_date", "تاريخ البداية", "Start date", False, "بصيغة تاريخ Excel أو YYYY-MM-DD", "date"),
    ("end_date", "تاريخ النهاية", "End date", False, "يجب أن يكون بعد تاريخ البداية", "date"),
    ("status", "الحالة", "Status", True, "اختر من القائمة", "status"),
    ("manager_name", "مدير المشروع", "Project manager", False, "", "text"),
    ("location", "الموقع", "Location", False, "", "text"),
    ("notes", "ملاحظات", "Notes", False, "", "text"),
]
statuses = "مخطط,قائم,متوقف مؤقتاً,مكتمل,مغلق"
dv_status = DataValidation(type="list", formula1=f'"{statuses}"', allow_blank=False, showErrorMessage=True, errorTitle="الحالة", error="اختر حالة من القائمة")
dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True, showErrorMessage=True, errorTitle="نسبة", error="النسبة يجب أن تكون بين 0 و 100")
dv_money = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle="مبلغ", error="المبلغ يجب أن يكون رقماً موجباً")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True, errorTitle="تاريخ", error="أدخل تاريخاً صحيحاً")
for dv in (dv_status, dv_pct, dv_money, dv_date): ws.add_data_validation(dv)
for i, (key, ar, en, req, note, kind) in enumerate(fields, start=5):
    ws[f"A{i}"] = ar + (" *" if req else ""); ws[f"B{i}"] = en; ws[f"D{i}"] = note
    ws[f"E{i}"] = key  # machine key read by the importer (hidden)
    for c in "ABCD": ws[f"{c}{i}"].border = box; ws[f"{c}{i}"].font = base; ws[f"{c}{i}"].alignment = right
    ws[f"A{i}"].font = bold
    if req: ws[f"A{i}"].fill = PatternFill("solid", fgColor=REQ)
    cell = ws[f"C{i}"]
    if kind == "money": dv_money.add(cell); cell.number_format = "#,##0.00"
    elif kind == "pct": dv_pct.add(cell); cell.number_format = "0.00"
    elif kind == "date": dv_date.add(cell); cell.number_format = "yyyy-mm-dd"
    elif kind == "status": dv_status.add(cell); cell.value = "قائم"
    elif kind == "int": dv_money.add(cell); cell.number_format = "0"
last = 4 + len(fields)
ws.column_dimensions["E"].hidden = True
for c, w in zip("ABCD", [42, 30, 34, 60]): ws.column_dimensions[c].width = w
r = last + 2
ws[f"A{r}"] = "إجمالي جدول الكميات (محسوب)"; ws[f"A{r}"].font = bold
ws[f"C{r}"] = "=SUM('جدول الكميات'!F:F)"; ws[f"C{r}"].number_format = "#,##0.00"; ws[f"C{r}"].font = bold
ws[f"A{r+1}"] = "الفرق عن قيمة العقد"; ws[f"A{r+1}"].font = bold
cv = 5 + [f[0] for f in fields].index("contract_value")
ws[f"C{r+1}"] = f"=C{cv}-C{r}"; ws[f"C{r+1}"].number_format = "#,##0.00"
ws[f"D{r+1}"] = f'=IF(ROUND(C{r+1},2)=0,"✔ مطابق","⚠ قيمة العقد لا تساوي إجمالي جدول الكميات")'
ws.conditional_formatting.add(f"D{r+1}", FormulaRule(formula=[f"ROUND(C{r+1},2)<>0"], font=Font(color="B00020", bold=True)))
ws.freeze_panes = "A5"

# ── Sheet 2: BOQ ────────────────────────────────────────────────────────
wq = wb.create_sheet("جدول الكميات"); wq.sheet_view.rightToLeft = True
heads = [("رقم البند", "Item no.", 12), ("الوصف *", "Description", 55), ("الوحدة", "Unit", 12),
         ("الكمية *", "Quantity", 14), ("سعر الوحدة *", "Unit price", 16), ("الإجمالي", "Amount", 18),
         ("الكمية المفوترة سابقاً", "Previously billed qty", 20)]
for j, (ar, en, w) in enumerate(heads, start=1):
    c = wq.cell(row=1, column=j, value=ar); c.font = hfont; c.fill = hfill; c.alignment = center; c.border = box
    e = wq.cell(row=2, column=j, value=en); e.font = Font(name="Arial", size=9, color="7A7A7A"); e.alignment = center; e.fill = PatternFill("solid", fgColor=GREY)
    wq.column_dimensions[chr(64 + j)].width = w
for j in (2, 4, 5): wq.cell(row=1, column=j).fill = PatternFill("solid", fgColor="8B1A1A")
dvq = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle="رقم", error="أدخل رقماً موجباً")
wq.add_data_validation(dvq)
samples = [("1", "أعمال الحفر والردم (مثال — احذفه)", "م3", 1200, 45),
           ("2", "خرسانة مسلحة للقواعد (مثال — احذفه)", "م3", 350, 780),
           ("3", "توريد وتركيب لوحات كهرباء (مثال — احذفه)", "عدد", 12, 6500)]
for i in range(3, 503):
    s = samples[i - 3] if i - 3 < len(samples) else None
    if s:
        for j, v in enumerate(s, start=1): wq.cell(row=i, column=j, value=v)
        for j in range(1, 8): wq.cell(row=i, column=j).font = Font(name="Arial", size=11, italic=True, color="8A8A8A")
    wq.cell(row=i, column=6, value=f'=IF(OR(D{i}="",E{i}=""),"",ROUND(D{i}*E{i},2))')
    for j in range(1, 8):
        c = wq.cell(row=i, column=j); c.border = box
        if j in (4, 7): c.number_format = "#,##0.000"
        if j in (5, 6): c.number_format = "#,##0.00"
    dvq.add(f"D{i}"); dvq.add(f"E{i}"); dvq.add(f"G{i}")
wq.conditional_formatting.add("G3:G502", FormulaRule(formula=['AND(G3<>"",D3<>"",G3>D3)'], fill=PatternFill("solid", fgColor=REQ), font=Font(color="B00020", bold=True)))
wq.conditional_formatting.add("B3:B502", FormulaRule(formula=['AND(B3="",OR(D3<>"",E3<>""))'], fill=PatternFill("solid", fgColor=REQ)))
wq.freeze_panes = "A3"

# ── Sheet 3: new customer (optional) ───────────────────────────────────
wc = wb.create_sheet("العميل"); wc.sheet_view.rightToLeft = True
wc["A1"] = "عبّئ هذه الورقة فقط إذا العميل غير موجود في النظام · Only if the customer is new"; wc["A1"].font = Font(name="Arial", size=12, bold=True)
cfields = [("name_ar", "اسم العميل (عربي) *", True), ("name_en", "اسم العميل (إنجليزي)", False),
           ("vat_number", "الرقم الضريبي (15 رقم يبدأ وينتهي بـ 3)", False), ("cr_number", "السجل التجاري (10 أرقام)", False),
           ("building_no", "رقم المبنى (4 أرقام)", False), ("street", "الشارع", False), ("district", "الحي", False),
           ("city", "المدينة", False), ("postal_code", "الرمز البريدي (5 أرقام)", False),
           ("email", "البريد الإلكتروني", False), ("phone", "الجوال", False)]
for c, h in zip("ABC", ["الحقل", "القيمة · Value", "key"]):
    wc[f"{c}3"] = h; wc[f"{c}3"].font = hfont; wc[f"{c}3"].fill = hfill; wc[f"{c}3"].alignment = center
for i, (k, ar, req) in enumerate(cfields, start=4):
    wc[f"A{i}"] = ar; wc[f"C{i}"] = k; wc[f"A{i}"].font = bold
    for col in "AB": wc[f"{col}{i}"].border = box; wc[f"{col}{i}"].alignment = right
    wc[f"B{i}"].number_format = "@"   # keep leading zeros in VAT/CR/postal code
    if req: wc[f"A{i}"].fill = PatternFill("solid", fgColor=REQ)
wc["A16"] = "ملاحظة: الرقم الضريبي والعنوان الوطني إلزاميان لإصدار فاتورة ضريبية (B2B) للعميل."; wc["A16"].font = Font(name="Arial", size=10, color="7A7A7A")
wc.column_dimensions["A"].width = 42; wc.column_dimensions["B"].width = 40; wc.column_dimensions["C"].hidden = True

# ── Sheet 4: instructions ──────────────────────────────────────────────
wi = wb.create_sheet("تعليمات"); wi.sheet_view.rightToLeft = True
lines = [
    ("قالب رفع مشروع إلى GDXora", Font(name="Arial", size=15, bold=True)),
    ("", None),
    ("1) ورقة «المشروع»: بيانات العقد. الحقول المعلّمة بـ * إلزامية.", base),
    ("2) ورقة «جدول الكميات»: بند في كل سطر. احذف أسطر الأمثلة الرمادية قبل الرفع.", base),
    ("   • الإجمالي يُحسب تلقائياً (الكمية × سعر الوحدة) — لا تكتب فيه.", base),
    ("   • «الكمية المفوترة سابقاً» للمشاريع القائمة فقط: الكمية التي صدر بها مستخلص قبل النظام. اتركها فارغة لمشروع جديد.", base),
    ("     النظام يبدأ المستخلص القادم من هذه الكمية، فلا تتكرر فوترتها.", base),
    ("   • للمشاريع القائمة عبّئ أيضاً في ورقة «المشروع»: عدد المستخلصات السابقة والدفعة المقدمة المستردة.", base),
    ("3) ورقة «العميل»: فقط إذا العميل غير مسجّل في النظام.", base),
    ("4) تأكد أن «قيمة العقد» تساوي إجمالي جدول الكميات (يظهر التحقق أسفل ورقة المشروع).", base),
    ("", None),
    ("عند الرفع: النظام يعرض معاينة كاملة، وأي خطأ يظهر بالأحمر مع رقم السطر وسبب الخطأ،", bold),
    ("ولا يُحفظ شيء حتى تكون كل البيانات سليمة وتضغط «تأكيد الاستيراد».", bold),
    ("", None),
    ("لا تغيّر أسماء الأوراق أو ترتيب الأعمدة.", Font(name="Arial", size=11, bold=True, color="B00020")),
]
for i, (txt, f) in enumerate(lines, start=1):
    wi[f"A{i}"] = txt
    if f: wi[f"A{i}"].font = f
wi.column_dimensions["A"].width = 120
wb.move_sheet("تعليمات", offset=-3); wb.active = 1
wb.save(out)
